from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DOCX = "梁博一-实习简历.docx"
LATIN_FONT = "Calibri"
CHINESE_FONT = "Hiragino Sans GB W3"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
INK = RGBColor(32, 32, 32)


def set_run_font(run, size=None, bold=False, color=None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), LATIN_FONT)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)

    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    r_pr.append(color_el)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)

    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), LATIN_FONT)
    font.set(qn("w:hAnsi"), LATIN_FONT)
    font.set(qn("w:eastAsia"), CHINESE_FONT)
    font.set(qn("w:cs"), LATIN_FONT)
    r_pr.append(font)

    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_contact_line(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=10)

    items = [
        ("城市：深圳", None),
        ("手机：18859047927", None),
        ("邮箱：byoneliang@gmail.com", "mailto:byoneliang@gmail.com"),
        ("GitHub：https://github.com/byoneliang-arch", "https://github.com/byoneliang-arch"),
        ("个人主页：https://byoneliang-arch.github.io/byoneliang.github.io/", "https://byoneliang-arch.github.io/byoneliang.github.io/"),
    ]

    for i, (text, url) in enumerate(items):
        if i:
            sep = p.add_run("  |  ")
            set_run_font(sep, 9.5, color=MUTED)
        if url:
            add_hyperlink(p, text, url)
        else:
            run = p.add_run(text)
            set_run_font(run, 9.5, color=MUTED)


def add_section(doc, title):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=10, after=4)
    run = p.add_run(title)
    set_run_font(run, 13, bold=True, color=BLUE)
    p.paragraph_format.keep_with_next = True

    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2F3")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=3, line=1.12)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    run = p.add_run(text)
    set_run_font(run, 9.8, color=INK)
    return p


def add_project(doc, name, links, tech, description, bullets, outcomes=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(name)
    set_run_font(r, 11.2, bold=True, color=DARK_BLUE)

    link_p = doc.add_paragraph()
    set_paragraph_spacing(link_p, after=2, line=1.05)
    run = link_p.add_run("链接：")
    set_run_font(run, 9.2, bold=True, color=MUTED)
    for i, (label, url) in enumerate(links):
        if i:
            sep = link_p.add_run("  |  ")
            set_run_font(sep, 9.2, color=MUTED)
        add_hyperlink(link_p, f"{label}：{url}", url)

    tech_p = doc.add_paragraph()
    set_paragraph_spacing(tech_p, after=2, line=1.05)
    run = tech_p.add_run("技术栈：")
    set_run_font(run, 9.2, bold=True, color=MUTED)
    run = tech_p.add_run(tech)
    set_run_font(run, 9.2, color=MUTED)

    desc_p = doc.add_paragraph()
    set_paragraph_spacing(desc_p, after=3, line=1.08)
    run = desc_p.add_run(description)
    set_run_font(run, 9.8, color=INK)

    for item in bullets:
        add_bullet(doc, item)

    if outcomes:
        for item in outcomes:
            add_bullet(doc, item)


def build():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), LATIN_FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, after=2)
    run = name.add_run("梁博一")
    set_run_font(run, 21, bold=True, color=RGBColor(11, 37, 69))

    target = doc.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(target, after=4)
    run = target.add_run("求职意向：AI 应用开发实习")
    set_run_font(run, 10.5, bold=True, color=DARK_BLUE)

    add_contact_line(doc)

    add_section(doc, "个人简介")
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=5, line=1.12)
    run = p.add_run(
        "深圳北理莫斯科大学数学与物理专业本科生，2024 级大二，预计 2028 年毕业。"
        "正在系统学习 Python 编程、AI 应用开发、RAG、Agent 开发与 GitHub 项目开发。"
        "已独立完成并部署 AI PDF 工具、Canvas 横版小游戏和数据可视化全栈应用，"
        "希望在实习中继续提升工程能力、AI 应用开发能力和产品落地能力。"
    )
    set_run_font(run, 9.8, color=INK)

    add_section(doc, "教育经历")
    edu = doc.add_paragraph()
    set_paragraph_spacing(edu, after=2)
    r = edu.add_run("深圳北理莫斯科大学")
    set_run_font(r, 10.4, bold=True, color=INK)
    r = edu.add_run("  |  数学与物理专业 本科  |  2024 级 大二，预计 2028 年毕业")
    set_run_font(r, 9.7, color=INK)
    add_bullet(doc, "相关基础：数学分析、高等代数、概率统计、物理建模、编程基础等。")
    add_bullet(doc, "重点方向：AI 应用开发、Python 编程、数据分析、Web 项目开发。")

    add_section(doc, "技能")
    skills = [
        "编程语言：Python、JavaScript、HTML、CSS",
        "AI 应用：AI 应用开发、RAG 基础、Agent 开发基础、Streamlit",
        "Web 开发：Vue 3、Vite、Flask、Canvas、响应式页面开发",
        "数据与可视化：ECharts、SQLite、CSV/Excel 数据处理、基础数据分析",
        "工程工具：Git、GitHub、Docker、Gunicorn、Render、GitHub Actions",
        "其他：机器学习基础、数学与物理建模思维",
    ]
    for item in skills:
        add_bullet(doc, item)

    add_section(doc, "项目经历")
    add_project(
        doc,
        "AI PDF 知识助手",
        [
            ("GitHub", "https://github.com/byoneliang-arch/ai-pdf-assistant"),
            ("在线演示", "https://ai-pdf-assistant-lby.streamlit.app/"),
        ],
        "Python、Streamlit",
        "支持上传 PDF 并进行智能检索的 AI 工具，面向文档阅读、资料查询和知识问答场景。",
        [
            "使用 Streamlit 搭建 Web 交互界面，实现 PDF 上传与检索结果展示。",
            "围绕 PDF 文档信息检索流程进行功能设计，提升用户查询资料的效率。",
            "将项目部署到线上演示环境，并通过 GitHub 管理项目代码。",
        ],
        ["熟悉了 Python Web 小工具从功能设计、代码实现到线上部署的完整流程。"],
    )

    add_project(
        doc,
        "小猪大冒险",
        [
            ("GitHub", "https://github.com/byoneliang-arch/pig-adventure1"),
            ("在线演示", "https://byoneliang-arch.github.io/pig-adventure1/"),
        ],
        "HTML、CSS、JavaScript、Canvas",
        "原创像素风 2D 横版网页小游戏，包含森林、暮色山地、雪山冰湖三关，以及金币收集、障碍规避、踩怪和终点判定。",
        [
            "使用 Canvas 实现 2D 游戏画面绘制、角色移动、关卡场景和基础碰撞逻辑。",
            "使用原生 HTML、CSS、JavaScript 完成游戏页面与交互逻辑，部署到 GitHub Pages。",
            "通过完整小游戏项目训练了交互设计、动画循环、状态管理和前端调试能力。",
        ],
    )

    add_project(
        doc,
        "本地数据记录与多折线图分析工具",
        [
            ("GitHub", "https://github.com/byoneliang-arch/line-chart-tool"),
            ("在线演示", "https://line-chart-tool.onrender.com"),
        ],
        "Vue 3、Vite、ECharts、Flask、SQLite、Docker、Gunicorn、Render、GitHub Actions",
        "用于记录和分析日期型数据的全栈 Web 应用，支持多条自定义折线、日期筛选、数据编辑、最大最小值标记、CSV/Excel 导出和图表导出。",
        [
            "使用 Vue 3 和 Vite 构建前端界面，实现数据录入、编辑、筛选和交互式图表展示。",
            "使用 ECharts 实现多折线趋势分析、最大值/最小值标记和图表导出能力。",
            "使用 Flask 与 SQLite 搭建后端接口和数据存储逻辑，并通过 Docker、Gunicorn、Render 完成部署。",
        ],
        ["完整实践了前后端分离项目的开发、部署和维护流程。"],
    )

    add_section(doc, "自我评价")
    for item in [
        "具备数学与物理背景，逻辑分析能力较强，适合处理 AI 应用、数据分析和工程问题。",
        "有持续完成实际项目的习惯，能够从需求出发独立完成开发、部署和迭代。",
        "对 AI 应用开发、AI 游戏开发、RAG 系统和 Agent 开发保持长期学习兴趣，希望在真实业务中继续提升工程能力。",
        "目前仍在持续开发 AI 项目、游戏项目和个人作品集项目，不断通过实际项目提升工程实践能力。",
    ]:
        add_bullet(doc, item)

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build()
